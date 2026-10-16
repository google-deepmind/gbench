# -*- coding: utf-8 -*-
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Canonical fidelity of the agentic tool suites: gaia, deepsearch_qa, mcp_atlas,
bfcl_v4_agentic and the per-suite tool-loop protocol in base.py.

Offline: the model endpoint, the judge cascade, the search backend and the dataset hubs are
all replaced at their seams.
"""

import asyncio
import base64
import csv
import hashlib
import io
import json
import zipfile
from pathlib import Path
from unittest import mock

import pytest

from gbench.runners.eval_suites import base

REPO = Path(__file__).resolve().parents[1]


# --------------------------------------------------------------------------- #
# base.py: per-suite tool-loop protocol
# --------------------------------------------------------------------------- #
_ASK = base.Reply(text="", tool_calls=[{"id": "c1", "function": {
    "name": "web_search", "arguments": '{"query":"birds"}'}}], finish_reason="tool_calls")


def _ask_with_content(content):
    calls = [{"id": "c1", "function": {"name": "web_search", "arguments": '{"query":"birds"}'}}]
    prefix = base._rendered_tool_call_prefix(calls)
    return base.Reply(text=prefix + content, tool_calls=calls, finish_reason="tool_calls")


def _run_loop(replies, executor, **kw):
    seen = {"n": 0, "convos": [], "max_tokens": []}

    async def _fake_send(**k):
        seen["convos"].append(list(k["messages"]))
        seen["max_tokens"].append(k["max_output_tokens"])
        r = replies[min(seen["n"], len(replies) - 1)]
        seen["n"] += 1
        return r

    with mock.patch.object(base, "_send_single_request", _fake_send):
        res = base.run_eval_suite(
            eval_name="w2_b4_loop", model_name="m", base_url="http://x", concurrency=1,
            samples=[([{"role": "user", "content": "q"}], "3", {})],
            eval_fn=lambda resp, g: g in (resp or ""), tool_executor=executor, **kw)
    return res, seen


async def _noop_tool(name, args):
    return '{"results": []}'


def test_house_loop_is_unchanged_without_a_protocol():
    res, seen = _run_loop([_ASK], _noop_tool, max_tool_rounds=2)
    assert seen["n"] == 4                                # 1 + 2 rounds + forced final
    assert res["tool_rounds"]["forced_final_answers"] == 1
    proto = res["tool_loop_protocol"]
    assert proto["reference"] == "gbench house loop" and proto["max_rounds"] == 2
    assert proto["house_turn_ceiling"] is True and proto["forced_final"] is True


def test_house_ceiling_still_caps_turns_and_a_protocol_can_opt_out(monkeypatch):
    base.set_run_knobs(max_output_tokens=65536)
    monkeypatch.delenv("GBENCH_TOOL_LOOP_MAX_OUTPUT_TOKENS", raising=False)
    answer = base.Reply(text="3", tool_calls=None, finish_reason="stop")
    _, seen = _run_loop([answer], _noop_tool)
    assert seen["max_tokens"][0] == base.TOOL_LOOP_MAX_OUTPUT_TOKENS
    proto = base.ToolLoopProtocol(reference="t", house_turn_ceiling=False)
    res, seen = _run_loop([answer], _noop_tool, tool_loop=proto)
    assert seen["max_tokens"][0] == 65536, "the operator's budget must not be cut per turn"
    assert res["tool_loop_protocol"]["per_turn_max_tokens"] == 65536


def test_no_forced_final_stops_at_the_cap_and_grades_the_last_content():
    proto = base.ToolLoopProtocol(reference="t", max_rounds=2, forced_final=False,
                                  rendered_calls_in_history=False)
    replies = [_ask_with_content("I found 3 birds so far."), _ASK, _ASK]
    res, seen = _run_loop(replies, _noop_tool, tool_loop=proto)
    assert seen["n"] == 2                                # 2 model turns, no extra request
    for convo in seen["convos"]:
        assert not any("No further tool calls" in str(m.get("content")) for m in convo)
    tr = res["sample_traces"][0]
    assert tr["response_text"] == "I found 3 birds so far."
    assert tr["tool_loop_end"] == "max_rounds"
    assert res["tool_rounds"]["loop_ends"] == {"max_rounds": 1}
    assert res["tool_rounds"]["forced_final_answers"] == 0


def test_tool_call_cap_stops_before_executing_past_it():
    ran = []

    async def ex(name, args):
        ran.append(name)
        return "{}"

    two_calls = base.Reply(text="", tool_calls=[
        {"id": "a", "function": {"name": "t1", "arguments": "{}"}},
        {"id": "b", "function": {"name": "t2", "arguments": "{}"}}], finish_reason="tool_calls")
    proto = base.ToolLoopProtocol(reference="t", max_rounds=256, max_tool_calls=3,
                                  forced_final=False)
    res, seen = _run_loop([two_calls], ex, tool_loop=proto)
    assert ran == ["t1", "t2", "t1"]
    assert seen["n"] == 2
    tr = res["sample_traces"][0]
    assert tr["tool_loop_end"] == "max_tool_calls" and tr["tool_calls_executed"] == 3


def test_uncapped_rounds_run_until_the_model_answers():
    proto = base.ToolLoopProtocol(reference="t", max_rounds=None, forced_final=False)
    answer = base.Reply(text="3", tool_calls=None, finish_reason="stop")
    res, seen = _run_loop([_ASK] * 12 + [answer], _noop_tool, tool_loop=proto)
    assert seen["n"] == 13
    assert res["sample_traces"][0]["tool_loop_end"] == "answered"
    assert res["correct_answers"] == 1


def test_no_forced_final_means_no_ran_out_of_room_turn():
    truncated = base.Reply(text="", tool_calls=None, finish_reason="length")
    proto = base.ToolLoopProtocol(reference="t", forced_final=False)
    _, seen = _run_loop([truncated], _noop_tool, tool_loop=proto)
    assert seen["n"] == 1
    _, seen = _run_loop([truncated, base.Reply(text="3", tool_calls=None)], _noop_tool)
    assert seen["n"] == 2, "the house loop still recovers"


def test_history_carries_the_models_own_content_not_the_rendering():
    proto = base.ToolLoopProtocol(reference="t", rendered_calls_in_history=False)
    answer = base.Reply(text="3", tool_calls=None, finish_reason="stop")
    _, seen = _run_loop([_ask_with_content("Let me look."), answer], _noop_tool, tool_loop=proto)
    assistant = [m for m in seen["convos"][1] if m["role"] == "assistant"][0]
    assert assistant["content"] == "Let me look."
    assert assistant["tool_calls"][0]["function"]["name"] == "web_search"
    _, seen = _run_loop([_ask_with_content("Let me look."), answer], _noop_tool)
    assistant = [m for m in seen["convos"][1] if m["role"] == "assistant"][0]
    assert assistant["content"].startswith("web_search("), "house loop unchanged"


def test_reply_content_matches_the_rendering_in_send_single_request():
    """`_rendered_tool_call_prefix` must stay byte-identical to `_send_single_request`."""
    body = {"choices": [{"message": {"content": "Checking.\nTwo lines.", "tool_calls": [
        {"id": "1", "type": "function",
         "function": {"name": "web_search", "arguments": '{\n  "query": "a, b"\n}'}},
        {"id": "2", "type": "function", "function": {"name": "f", "arguments": "not json"}}]},
        "finish_reason": "tool_calls"}], "usage": {"completion_tokens": 5}}

    class _Resp:
        status = 200
        async def json(self): return body
        async def text(self): return ""
        async def __aenter__(self): return self
        async def __aexit__(self, *a): return False

    class _Session:
        def post(self, *a, **k): return _Resp()

    reply = asyncio.run(base._send_single_request(
        session=_Session(), api_url="u", model_name="m", messages=[], extra_payload={},
        semaphore=asyncio.Semaphore(1), pbar=None))
    assert reply.tool_calls and base.reply_content(reply) == "Checking.\nTwo lines."


# --------------------------------------------------------------------------- #
# deepsearch_qa: Google DeepMind DeepSearchQA
# --------------------------------------------------------------------------- #
from gbench.runners.eval_suites import deepsearch_qa as DSQA  # noqa: E402


def test_dsqa_rater_prompt_is_the_notebook_text_verbatim():
    """sha256 of DEEPSEARCH_QA_PROMPT / GRADER_RATING_OUTPUT_EXAMPLE as defined in the Kaggle
    starter notebook (andrewmingwang/deepsearchqa-starter-code v1)."""
    assert hashlib.sha256(DSQA.DEEPSEARCH_QA_PROMPT.encode()).hexdigest() == \
        "16961a857c49a523f1e31c45f953885204e31f9892aeb6ddc4af3b7a0a8c92fc"
    assert hashlib.sha256(DSQA.GRADER_RATING_OUTPUT_EXAMPLE.encode()).hexdigest() == \
        "077b7deb37edd1752415aec7781a0a66d32daf3e235a68eca3585c0ee0d5df0f"
    p = DSQA.grader_prompt(" Q? ", "Set Answer", " A, B ", " resp ")
    assert p.startswith(DSQA.DEEPSEARCH_QA_PROMPT)
    assert "<prompt>\nQ?\n</prompt>" in p and "Prompt Type: Set Answer" in p
    assert "<answer>\nA, B\n</answer>" in p and p.endswith("Rating:")


def test_dsqa_parse_rating_follows_the_notebook_rules():
    fenced = ('text ```json\n{"Answer Correctness": {"Explanation": "x", '
              '"Correctness Details": {"A": true, "B": false}, "Excessive Answers": ["C"]}}\n```')
    details, excessive, err = DSQA.parse_rating(fenced)
    assert err is None and details == {"A": True, "B": False} and excessive == ["C"]
    no_excess = '{"Answer Correctness": {"Explanation": "x", "Correctness Details": {"A": true}}}'
    assert DSQA.parse_rating(no_excess)[1] == []           # absent key is valid
    assert DSQA.parse_rating("not json")[2]
    assert DSQA.parse_rating('{"Answer Correctness": {"Explanation": "x", '
                             '"Correctness Details": {"A": "yes"}}}')[2]
    assert DSQA.parse_rating("")[2]


@pytest.mark.parametrize("details,excessive,category,f1", [
    ({"A": True, "B": True}, [], "fully_correct", 1.0),
    ({"A": True, "B": True}, ["C"], "correct_with_extraneous_answers", 0.8),
    ({"A": True, "B": False}, [], "partially_correct", 2 * 1.0 * 0.5 / 1.5),
    ({"A": False}, ["X"], "fully_incorrect", 0.0),
    ({}, [], "fully_correct", 0.0),
])
def test_dsqa_item_scoring(details, excessive, category, f1):
    item = DSQA.score_item(details, excessive)
    assert item["category"] == category
    assert item["f1_score"] == pytest.approx(f1)
    assert item["is_all_correct"] is (category == "fully_correct")


def _dsqa_trace(resp, answer_type="Set Answer", gold="A, B"):
    return {"response_text": resp, "gold_answer": gold,
            "messages": [{"role": "user", "content": "Which?"}],
            "extra_payload": {"category": "Science", "answer_type": answer_type}}


def test_dsqa_judge_grades_and_summarises(monkeypatch):
    replies = iter([
        '{"Answer Correctness": {"Explanation": "x", "Correctness Details": '
        '{"A": true, "B": true}, "Excessive Answers": []}}',
        '{"Answer Correctness": {"Explanation": "x", "Correctness Details": '
        '{"A": true, "B": false}, "Excessive Answers": ["C"]}}',
        "garbage",
    ])
    prompts = []

    async def _judge(prompt, **kw):
        prompts.append(prompt)
        return next(replies), "gemini-3.6-flash"

    monkeypatch.setattr(DSQA, "judge_generate_cascade", _judge)
    traces = [_dsqa_trace("A and B"), _dsqa_trace("A and C"), _dsqa_trace("?"), _dsqa_trace("")]
    asyncio.run(DSQA._async_judge_deepsearch_qa(traces))
    assert len(prompts) == 3, "an empty response is not sent to the rater"
    assert "Prompt Type: Set Answer" in prompts[0]
    assert traces[0]["is_correct"] is True and traces[1]["is_correct"] is False
    assert traces[2]["judge_grade"] == "unparsed" and traces[2]["scoring_excluded"] is True
    assert traces[3]["dsqa_empty_response"] is True and traces[3]["is_correct"] is False
    s = DSQA.summarize(traces)
    assert s["num_rated"] == 3 and s["num_invalid_auto_rater_response"] == 1
    assert s["fully_correct"] == pytest.approx(33.33)
    assert s["fully_correct_excluding_empty"] == 50.0
    # per-item F1 averaged: 1.0, 0.5 (P=1/2, R=1/2), 0.0 for the empty response
    assert s["f1_score"] == pytest.approx(round(100 * (1.0 + 0.5 + 0.0) / 3, 2))


def test_dsqa_judge_outage_is_excluded(monkeypatch):
    async def _judge(prompt, **kw):
        return None, "JUDGE_OUTAGE"
    monkeypatch.setattr(DSQA, "judge_generate_cascade", _judge)
    traces = [_dsqa_trace("A")]
    asyncio.run(DSQA._async_judge_deepsearch_qa(traces))
    assert traces[0]["judge_grade"] == "JUDGE_OUTAGE"
    assert DSQA.summarize(traces)["num_rated"] == 0


def test_dsqa_loader_sends_the_raw_problem_and_keeps_literal_none(tmp_path, monkeypatch):
    path = tmp_path / "DSQA-full.csv"
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["problem", "problem_category", "answer", "answer_type"])
        w.writerow(["List the games.", "Media & Entertainment", "None", "Set Answer"])
        w.writerow(["Who won?", "Sports", "Brazil", "Single Answer"])
    seen = {}

    def _dl(**kw):
        seen.update(kw)
        return str(path)

    monkeypatch.setattr("huggingface_hub.hf_hub_download", _dl)
    samples = DSQA._load_deepsearch_qa_samples()
    assert seen["revision"] == DSQA.DATASET_REVISION and seen["repo_id"] == "google/deepsearchqa"
    msgs, gold, meta = samples[0]
    assert msgs == [{"role": "user", "content": "List the games."}]
    assert gold == "None", "pandas would turn the literal answer 'None' into NaN"
    assert meta == {"category": "Media & Entertainment", "answer_type": "Set Answer"}
    assert "Set Answer" not in json.dumps(msgs), "answer_type must not reach the model"


def test_dsqa_run_uses_the_reference_loop_and_flags_the_judge(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "k")
    monkeypatch.setattr(DSQA, "gemini_key_live_valid", lambda k: (True, ""))
    monkeypatch.setattr(DSQA, "_load_deepsearch_qa_samples",
                        lambda limit=None: [([{"role": "user", "content": "q"}], "A",
                                             {"category": "Science", "answer_type": "Single Answer"})])
    got = {}

    def _fake_run(**kw):
        got.update(kw)
        return {"sample_traces": [{"judge_model": "gemini-3.6-flash", "is_correct": True,
                                   "dsqa": {"category": "fully_correct", "precision": 1.0,
                                            "recall": 1.0, "f1_score": 1.0},
                                   "extra_payload": {"answer_type": "Single Answer"}}],
                "accuracy": 100.0}

    monkeypatch.setattr(DSQA, "run_eval_suite", _fake_run)
    res = DSQA.run_deepsearch_qa("m", "http://x", concurrency=1)
    assert got["tool_loop"] is DSQA.DSQA_TOOL_LOOP
    assert DSQA.DSQA_TOOL_LOOP.max_rounds is None and DSQA.DSQA_TOOL_LOOP.forced_final is False
    assert DSQA.DSQA_TOOL_LOOP.house_turn_ceiling is False
    assert got["extra_payload"]["tools"][0]["function"]["name"] == "web_search"
    assert res["fully_correct"] == 100.0 and res["f1_score"] == 100.0
    assert res["leaderboard_comparable"] is False
    assert "gemini-2.5-flash" in res["leaderboard_comparable_reason"]
    assert res["judge_models"] == {"gemini-3.6-flash": 1}


# --------------------------------------------------------------------------- #
# gaia: attachments
# --------------------------------------------------------------------------- #
from gbench.runners.eval_suites import gaia as GAIA  # noqa: E402


def _docx_bytes():
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           '<w:body><w:p><w:r><w:t>Gift list</w:t></w:r></w:p>'
           '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Harry</w:t></w:r></w:p></w:tc>'
           '<w:tc><w:p><w:r><w:t>Miguel</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
           '</w:body></w:document>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    return buf.getvalue()


def _pptx_bytes():
    slide = ('<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
             'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree>'
             '<p:sp><p:txBody><a:p><a:r><a:t>{}</a:t></a:r></a:p></p:txBody></p:sp>'
             '</p:spTree></p:cSld></p:sld>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("ppt/slides/slide10.xml", slide.format("crayfish"))
        z.writestr("ppt/slides/slide2.xml", slide.format("eels"))
    return buf.getvalue()


def test_gaia_images_and_audio_are_media_parts():
    png = GAIA._file_parts("a.png", b"\x89PNG....")
    assert png[1]["type"] == "image_url"
    assert png[1]["image_url"]["url"].startswith("data:image/png;base64,")
    jpg = GAIA._file_parts("a.jpg", b"\xff\xd8")
    assert jpg[1]["image_url"]["url"].startswith("data:image/jpeg;base64,")
    mp3 = GAIA._file_parts("a.mp3", b"ID3...")
    assert mp3[1] == {"type": "input_audio", "input_audio": {
        "data": base64.b64encode(b"ID3...").decode(), "format": "mp3"}}
    assert png[0]["text"] == "\n\nAttached file: a.png"


def test_gaia_documents_are_rendered_to_text():
    assert "Harry | Miguel" in GAIA._file_parts("x.docx", _docx_bytes())[0]["text"]
    pptx = GAIA._file_parts("x.pptx", _pptx_bytes())[0]["text"]
    assert pptx.index("Slide 1:\neels") < pptx.index("Slide 2:\ncrayfish"), "numeric slide order"
    csv_part = GAIA._file_parts("x.csv", b"a,b\n1,2\n")[0]["text"]
    assert csv_part == "\n\nAttached file: x.csv\na,b\n1,2"


def test_gaia_xlsx_keeps_cell_colors(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "START"
    ws["B1"].fill = PatternFill("solid", fgColor="FF00FF00")
    p = tmp_path / "map.xlsx"
    wb.save(p)
    text = GAIA.attachment_parts(str(p), "map.xlsx")[0]["text"]
    assert "START," in text and "FF00FF00: B1" in text


def test_gaia_zip_expands_to_its_members_and_long_text_is_clipped():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("notes.txt", "x" * (GAIA.MAX_ATTACHMENT_CHARS + 50))
        z.writestr("pic.png", b"\x89PNG")
    parts = GAIA._file_parts("bundle.zip", buf.getvalue())
    assert "2 file(s)" in parts[0]["text"]
    assert "truncated at 100000 characters" in parts[1]["text"]
    assert parts[3]["type"] == "image_url"


def test_gaia_unknown_attachment_type_is_a_hard_error():
    with pytest.raises(RuntimeError, match="no renderer"):
        GAIA._file_parts("model.step", b"...")


def _gaia_snapshot(tmp_path, file_bytes=b"a,b\n1,2\n"):
    pd = pytest.importorskip("pandas")
    vdir = tmp_path / "2023" / "validation"
    vdir.mkdir(parents=True)
    rows = [{"task_id": "t1", "Question": "Sum column b.", "Level": "1", "Final answer": "2",
             "file_name": "t1.csv", "file_path": "2023/validation/t1.csv"},
            {"task_id": "t2", "Question": "Capital of France?", "Level": "2",
             "Final answer": "Paris", "file_name": "", "file_path": ""}]
    pd.DataFrame(rows).to_parquet(vdir / "metadata.parquet")
    if file_bytes is not None:
        (vdir / "t1.csv").write_bytes(file_bytes)
    return str(tmp_path)


def test_gaia_loader_pins_the_revision_and_inlines_the_file(tmp_path, monkeypatch):
    snap = _gaia_snapshot(tmp_path)
    seen = {}

    def _snap(**kw):
        seen.update(kw)
        return snap

    monkeypatch.setattr("huggingface_hub.snapshot_download", _snap)
    samples = GAIA._load_gaia_samples()
    assert seen["revision"] == GAIA.DATASET_REVISION
    assert seen["allow_patterns"] == ["2023/validation/*"]
    by_q = {json.dumps(s[0][1]["content"])[:30]: s for s in samples}
    with_file = next(s for s in samples if s[2]["has_file"])
    assert with_file[0][0]["content"] == GAIA._GAIA_SYSTEM
    content = with_file[0][1]["content"]
    assert content[0] == {"type": "text", "text": "Sum column b."}
    assert "Attached file: t1.csv\na,b\n1,2" in content[1]["text"]
    assert with_file[2]["attachment_type"] == "csv"
    free = next(s for s in samples if not s[2]["has_file"])
    assert free[0][1]["content"] == "Capital of France?"
    assert len(by_q) == 2


def test_gaia_missing_attachment_or_renderer_dependency_hard_errors(tmp_path, monkeypatch):
    snap = _gaia_snapshot(tmp_path, file_bytes=None)
    monkeypatch.setattr("huggingface_hub.snapshot_download", lambda **kw: snap)
    with pytest.raises(RuntimeError, match="missing from the snapshot"):
        GAIA._load_gaia_samples()

    snap2 = _gaia_snapshot(tmp_path / "b")
    monkeypatch.setattr("huggingface_hub.snapshot_download", lambda **kw: snap2)

    def _no_dep(path, name):
        raise GAIA.AttachmentDependencyError("xlrd (for legacy .xls attachments)")

    monkeypatch.setattr(GAIA, "attachment_parts", _no_dep)
    with pytest.raises(RuntimeError, match="INFRA REQUIRED.*xlrd"):
        GAIA._load_gaia_samples()


def test_gaia_run_uses_its_reference_loop(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "k")
    monkeypatch.setattr(GAIA, "check_gaia_prerequisites", lambda: (True, ""))
    monkeypatch.setattr(GAIA, "gemini_key_live_valid", lambda key: (True, ""))
    monkeypatch.setattr(GAIA, "_load_gaia_samples", lambda limit=None: [])
    got = {}

    def _fake_run(**kw):
        got.update(kw)
        return {"sample_traces": [
            {"is_correct": True, "extra_payload": {"has_file": True, "attachment_type": "png"}},
            {"is_correct": False, "extra_payload": {"has_file": False}}]}

    monkeypatch.setattr(GAIA, "run_eval_suite", _fake_run)
    res = GAIA.run_gaia("m", "http://x", concurrency=1)
    assert got["tool_loop"] is GAIA.GAIA_TOOL_LOOP
    assert GAIA.GAIA_TOOL_LOOP.max_rounds == 20 and GAIA.GAIA_TOOL_LOOP.forced_final is True
    assert GAIA.GAIA_TOOL_LOOP.house_turn_ceiling is False
    assert res["accuracy_with_file"] == 100.0 and res["accuracy_file_free"] == 0.0
    assert res["accuracy_by_attachment_type"] == {"png": {"n": 1, "accuracy": 100.0}}
    assert res["leaderboard_comparable"] is False and res["leaderboard_comparable_reason"]


# --------------------------------------------------------------------------- #
# mcp_atlas
# --------------------------------------------------------------------------- #
from gbench.runners.eval_suites import mcp_atlas as MCP  # noqa: E402


def test_mcp_atlas_judge_prompt_is_upstream_verbatim():
    """sha256 of `_get_single_claim_evaluation_prompt` at scaleapi/mcp-atlas@f24ba3f (with the
    plus-minus sign), and no format instruction appended: the schema carries the format."""
    assert hashlib.sha256(MCP._CLAIM_JUDGE_PROMPT.encode()).hexdigest() == \
        "87e98017c03361bee3f792c0e92fdeb2484abaa04131bd1b2f8477ac86e6d434"
    assert "\u00b11 percentage points" in MCP._CLAIM_JUDGE_PROMPT
    assert "Reply with ONLY" not in MCP._CLAIM_JUDGE_PROMPT
    assert MCP._CLAIM_JUDGE_SCHEMA["properties"]["coverage_outcome"]["enum"] == [
        "fulfilled", "partially_fulfilled", "not_fulfilled"]


def test_mcp_atlas_claims_are_cleaned_like_score_claims():
    blob = json.dumps(["\u2022 The repo was created in \u201c2013\u201d.", "abc",
                       "2) Domain registered in 2006 \u2014 per WHOIS."])
    assert MCP.extract_claims(blob) == ['The repo was created in "2013',
                                        "Domain registered in 2006 - per WHOIS."]
    assert MCP.extract_claims("first claim here; second claim here") == [
        "first claim here", "second claim here"]


def test_mcp_atlas_loader_sends_the_raw_prompt(monkeypatch):
    rows = [{"TASK": "t1", "PROMPT": "Find the year gap.",
             "ENABLED_TOOLS": '["fetch_fetch","whois_whois_domain"]',
             "GTFA_CLAIMS": '["The repo was created in 2013."]'}]
    seen = {}

    def _ld(*a, **kw):
        seen.update(kw)
        return rows

    monkeypatch.setattr("datasets.load_dataset", _ld)
    # The tools come from the agent environment's /list-tools catalog, not from name stubs.
    catalog = [{"name": n, "description": f"{n} tool", "inputSchema": {"type": "object"}}
               for n in ("fetch_fetch", "whois_whois_domain")]
    (msgs, gold, meta), = MCP._load_mcp_atlas_samples(catalog=catalog)
    assert seen["revision"] == MCP.DATASET_REVISION
    assert msgs == [{"role": "user", "content": "Find the year gap."}]
    assert [t["function"]["name"] for t in meta["tools"]] == ["fetch_fetch", "whois_whois_domain"]
    rows[0]["GTFA_CLAIMS"] = "[]"
    with pytest.raises(RuntimeError, match="refusing to drop it"):
        MCP._load_mcp_atlas_samples(catalog=catalog)


def test_mcp_atlas_judge_uses_structured_output_and_the_075_pass_bar(monkeypatch):
    seen = []
    outcomes = iter(["fulfilled", "fulfilled", "fulfilled", "partially_fulfilled"])

    async def _judge(prompt, config=None):
        seen.append((prompt, config))
        return json.dumps({"claim_text": "c", "coverage_outcome": next(outcomes),
                           "justification": "j", "confidence_level": 0.9}), "gemini-3.6-flash"

    monkeypatch.setattr(MCP, "judge_generate_cascade", _judge)
    claims = json.dumps([f"claim number {i}" for i in range(4)])
    long_answer = "y" * (MCP._MAX_RESPONSE_CHARS + 10)
    tr = [{"gold_answer": claims, "response_text": long_answer}]
    asyncio.run(MCP._async_judge_mcp_atlas(tr))
    assert tr[0]["claim_recall"] == 0.875 and tr[0]["is_correct"] is True
    prompt, cfg = seen[0]
    assert prompt.endswith(MCP._TRUNCATION_NOTE + "\nINSTRUCTIONS:" +
                           MCP._CLAIM_JUDGE_PROMPT.split("\nINSTRUCTIONS:", 1)[1])
    assert cfg is None or getattr(cfg, "response_json_schema", None) == MCP._CLAIM_JUDGE_SCHEMA
    assert tr[0]["judge_models"] == ["gemini-3.6-flash"]


def test_mcp_atlas_run_headlines_the_leaderboard_pass_rate(monkeypatch):
    monkeypatch.setattr(MCP, "gemini_required_skip", lambda *a: None)
    monkeypatch.setattr(MCP, "check_mcp_atlas_prerequisites", lambda: (True, ""))
    monkeypatch.setattr(MCP, "_open_sandbox", lambda: MCP._Sandbox("http://env"))
    monkeypatch.setattr(MCP, "_wait_ready", lambda sb, t: None)
    monkeypatch.setattr(MCP, "_preflight", lambda sb: {"health_check": "passed"})
    monkeypatch.setattr(MCP, "_list_tools", lambda sb: [{"name": "fetch_fetch"}])
    monkeypatch.setattr(MCP, "_load_mcp_atlas_samples", lambda limit=None, catalog=None: [])
    got = {}

    def _fake_run(**kw):
        got.update(kw)
        tools = [{"function": {"name": "fetch_fetch"}}, {"function": {"name": "notion_x"}}]
        return {"sample_traces": [
            {"claim_recall": 0.8, "is_correct": True, "claim_scores": [1, 0.5],
             "judge_models": ["gemini-3.6-flash"], "extra_payload": {"tools": tools}},
            {"claim_recall": 0.6, "is_correct": False, "claim_scores": [1, 0.5, 0],
             "judge_models": ["gemini-3.6-flash"], "extra_payload": {"tools": tools}}],
            "accuracy": 50.0}

    monkeypatch.setattr(MCP, "run_eval_suite", _fake_run)
    res = MCP.run_mcp_atlas("m", "http://x", concurrency=1)
    assert got["tool_loop"] is MCP.MCP_ATLAS_TOOL_LOOP
    assert got["tool_executor"] is MCP.execute_mcp_tool
    assert MCP.MCP_ATLAS_TOOL_LOOP.max_tool_calls == 100 and MCP.MCP_ATLAS_TOOL_LOOP.max_rounds == 256
    assert MCP.MCP_ATLAS_TOOL_LOOP.forced_final is False
    assert res["accuracy"] == res["pass_rate_0.75"] == 50.0
    assert res["pass_rate_0.50"] == 100.0 and res["mean_coverage"] == 70.0
    assert res["leaderboard_comparable"] is False
    assert res["leaderboard_column"] == "Pass Rate % (Public 500)"
    assert "gemini-3.1-pro-preview" in res["leaderboard_comparable_reason"]
    assert "closed_book" not in res


# --------------------------------------------------------------------------- #
# bfcl_v4_agentic
# --------------------------------------------------------------------------- #
from gbench.runners.eval_suites import bfcl_v4_agentic as BFCL  # noqa: E402

ALL_CATS = ["memory_kv", "memory_vector", "memory_rec_sum",
            "web_search_base", "web_search_no_snippet"]


def test_bfcl_agentic_score_is_the_two_level_unweighted_mean():
    per_cat = {"memory_kv": {"correct": 31, "total": 155, "accuracy": 20.0},
               "memory_vector": {"correct": 31, "total": 155, "accuracy": 20.0},
               "memory_rec_sum": {"correct": 31, "total": 155, "accuracy": 20.0},
               "web_search_base": {"correct": 80, "total": 100, "accuracy": 80.0},
               "web_search_no_snippet": {"correct": 60, "total": 100, "accuracy": 60.0}}
    s = BFCL.agentic_scores(per_cat)
    assert s == {"summary_web_search": 70.0, "summary_memory": 20.0,
                 "overall_accuracy_agentic": 45.0}
    pooled = 100.0 * (31 * 3 + 80 + 60) / (155 * 3 + 200)
    assert abs(pooled - 45.0) > 5, "the pooled rate weighs memory 70% and must differ"


def test_bfcl_category_selection():
    assert BFCL.select_categories(ALL_CATS, None) == (ALL_CATS, False)
    cats, subset = BFCL.select_categories(ALL_CATS, "memory_kv, memory_vector,memory_rec_sum")
    assert cats == ["memory_kv", "memory_vector", "memory_rec_sum"] and subset is True
    with pytest.raises(ValueError, match="unknown"):
        BFCL.select_categories(ALL_CATS, "memory_kv,web_search")


def _bfcl_harness(monkeypatch, tmp_path, per_cat=None, backend=(BFCL.SEARCH_BACKEND, "")):
    calls = []
    monkeypatch.setenv("BFCL_PROJECT_ROOT", str(tmp_path))
    monkeypatch.setattr(BFCL, "check_bfcl_prerequisites", lambda url: (True, "", list(ALL_CATS)))
    monkeypatch.setattr(BFCL, "search_backend", lambda: backend)
    monkeypatch.setattr(BFCL, "_register_model", lambda m: m)
    monkeypatch.setattr(BFCL, "_parse_score_traces", lambda *a, **k: [])
    monkeypatch.setattr(BFCL, "_run_cli", lambda argv, env, t, model=None, backend=None:
                        (calls.append(list(argv)), (0, "ok"))[1])
    monkeypatch.setattr(BFCL, "_parse_scores", lambda root, model, cats: (
        per_cat if per_cat is not None
        else {c: {"correct": 1, "total": 2, "accuracy": 50.0} for c in cats}))
    return calls


def test_bfcl_full_gemini_search_run_headlines_the_agentic_score(monkeypatch, tmp_path):
    _bfcl_harness(monkeypatch, tmp_path)
    res = BFCL.run_bfcl_v4_agentic("m", "http://x", concurrency=1)
    assert res["accuracy"] == res["overall_accuracy_agentic"] == 50.0
    assert res["accuracy_pooled"] == 50.0
    # the deviations: the web search is served by Gemini grounding, and the no-think 0.0
    assert res["leaderboard_comparable"] is False
    assert res["leaderboard_comparable_reason"] == (
        BFCL.SEARCH_SUBSTITUTION_REASON
        + "; temperature 0.0 set by default (BFCL generates at 0.001)")


def test_bfcl_missing_backend_reports_web_search_and_explicit_memory_only_runs(monkeypatch,
                                                                            tmp_path):
    calls = _bfcl_harness(monkeypatch, tmp_path, backend=(None, "no web-search backend"))
    # the suite's own full run hard-errors (its score weighs the Web Search setting); bfcl's
    # agentic part reports the setting and runs Memory
    with pytest.raises(RuntimeError, match="INFRA REQUIRED"):
        BFCL.run_bfcl_v4_agentic("m", "http://x", concurrency=1)
    assert not [c for c in calls if c[0] == "generate"]
    with BFCL.suite_context("bfcl", "docs/evals/bfcl.md"):
        res = BFCL.run_bfcl_v4_agentic("m", "http://x", concurrency=1)
    gen = next(c for c in calls if c[0] == "generate")
    assert gen[gen.index("--test-category") + 1] == "memory_kv,memory_vector,memory_rec_sum"
    web = res["settings"]["web_search"]
    assert web["status"] == "infra_required" and "no web-search backend" in web["infra_required"]
    assert res["summary_web_search"] is None and res["overall_accuracy_agentic"] is None
    assert res["accuracy"] == res["summary_memory"] == 50.0
    assert res["status"] == "completed_with_errors" and res["leaderboard_comparable"] is False
    calls.clear()
    res = BFCL.run_bfcl_v4_agentic("m", "http://x", concurrency=1,
                                   eval_categories="memory_kv,memory_vector,memory_rec_sum")
    gen = next(c for c in calls if c[0] == "generate")
    assert gen[gen.index("--test-category") + 1] == "memory_kv,memory_vector,memory_rec_sum"
    assert res["summary_web_search"] is None and res["summary_memory"] == 50.0
    assert res["bfcl_report"]["categories_subset"] is True
    assert res["leaderboard_comparable"] is False
    assert "--eval-categories subset" in res["leaderboard_comparable_reason"]


def test_bfcl_unscored_category_is_not_reported_as_a_partial(monkeypatch, tmp_path):
    _bfcl_harness(monkeypatch, tmp_path, per_cat={
        "memory_kv": {"correct": 1, "total": 2, "accuracy": 50.0}})
    with pytest.raises(RuntimeError, match="produced no score"):
        BFCL.run_bfcl_v4_agentic("m", "http://x", concurrency=1)


def test_bfcl_eval_limit_id_file_failure_is_a_hard_error(monkeypatch, tmp_path):
    _bfcl_harness(monkeypatch, tmp_path)

    def _boom(*a, **k):
        raise KeyError("dataset")

    monkeypatch.setattr(BFCL, "_select_limited_ids", _boom)
    # the id file is built in the harness image; here in this process
    monkeypatch.setattr(BFCL, "run_in_image",
                        lambda name, module, func, kwargs, **k: getattr(BFCL, func)(**kwargs))
    with pytest.raises(RuntimeError, match="could not build the id file"):
        BFCL.run_bfcl_v4_agentic("m", "http://x", concurrency=1, limit=10)


# --------------------------------------------------------------------------- #
# docs match the code
# --------------------------------------------------------------------------- #
def _doc(name):
    return (REPO / "docs" / "evals" / f"{name}.md").read_text(encoding="utf-8")


def test_docs_pin_the_same_revisions_as_the_code():
    assert GAIA.DATASET_REVISION in _doc("gaia")
    assert DSQA.DATASET_REVISION in _doc("deepsearch_qa")
    assert MCP.DATASET_REVISION in _doc("mcp_atlas")


def test_gaia_doc_quotes_the_injected_turns_verbatim():
    import inspect
    src = inspect.getsource(base)
    doc = " ".join(_doc("gaia").split())
    for text in (
            "You ran out of room to think. Do not reason further. State your best answer now "
            "from what you already have, in the requested format.",
            "No further tool calls are possible. Answer now, using only what you already have, "
            "and state your final answer explicitly.",
            "That response could not be read because it was a tool call. Tools are disabled and "
            "no tool call can succeed. Reply with plain text only - no function call, no tool "
            "syntax - stating your final answer."):
        assert text in doc
        assert " ".join(text.split()[:6]) in " ".join(src.split())


@pytest.mark.parametrize("name", ["gaia", "deepsearch_qa", "mcp_atlas", "bfcl_v4_agentic"])
def test_suite_docs_are_plain_ascii(name):
    bad = [line for line in _doc(name).splitlines() if any(ord(c) > 127 for c in line)]
    assert not bad, bad[:3]


# --------------------------------------------------------------------------- #
# bfcl_v4_agentic category routing and stale score files
# --------------------------------------------------------------------------- #
def test_bfcl_agentic_categories_are_the_ones_the_cli_routes():
    from bfcl_eval.constants.category_mapping import TEST_COLLECTION_MAPPING
    from gbench.runners.evals import eval_category_names, route_eval_categories
    assert set(BFCL.AGENTIC_CATEGORIES) == set(TEST_COLLECTION_MAPPING["agentic"])
    assert eval_category_names()["bfcl_v4_agentic"] == set(BFCL.AGENTIC_CATEGORIES)
    # `--evals all --eval-categories simple_python`: the agentic suite gets nothing and runs
    # its full track; a mixed list gives it only its own names.
    assert route_eval_categories("simple_python", "bfcl_v4_agentic") is None
    assert route_eval_categories("simple_python,memory_kv", "bfcl_v4_agentic") == "memory_kv"


def test_bfcl_all_run_with_bfcl_categories_runs_the_full_agentic_track(monkeypatch, tmp_path):
    from gbench.runners.evals import route_eval_categories
    calls = _bfcl_harness(monkeypatch, tmp_path)
    res = BFCL.run_bfcl_v4_agentic(
        "m", "http://x", concurrency=1,
        eval_categories=route_eval_categories("simple_python", "bfcl_v4_agentic"))
    gen = next(c for c in calls if c[0] == "generate")
    assert gen[gen.index("--test-category") + 1] == ",".join(BFCL.AGENTIC_CATEGORIES)
    assert res["bfcl_report"]["categories_subset"] is False
    with pytest.raises(ValueError, match="unknown --eval-categories"):   # typos still error
        BFCL.run_bfcl_v4_agentic("m", "http://x", concurrency=1, eval_categories="memory_kvv")


def test_bfcl_harness_with_other_agentic_categories_is_refused(monkeypatch, tmp_path):
    monkeypatch.setattr(BFCL, "_agentic_categories", lambda: ["memory_kv", "web_search_v2"])
    # the mapping is read in the harness image (its probe, run here in this process)
    monkeypatch.setenv("GBENCH_BFCL_PROJECT_ROOT", str(tmp_path))
    monkeypatch.setattr(BFCL.urllib.request, "urlopen", lambda *a, **k: _Ok())
    monkeypatch.setattr(BFCL, "image_problem", lambda: None)
    monkeypatch.setattr(BFCL, "run_in_image", lambda *a, **k: {
        **BFCL.image_probe(), "release": BFCL.BFCL_EVAL_RELEASE, "agentic_mismatch": None})
    ok, reason, cats = BFCL.check_bfcl_prerequisites("http://x")
    assert ok is False and cats == [] and "web_search_v2" in reason


class _Ok:
    status = 200

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def _write_scores(root, model, cats, acc=0.9, count=10):
    d = root / "score" / model.replace("/", "_") / "agentic"
    d.mkdir(parents=True, exist_ok=True)
    for c in cats:
        (d / f"BFCL_v4_{c}_score.json").write_text(
            json.dumps({"accuracy": acc, "correct_count": round(acc * count),
                        "total_count": count}) + "\n", encoding="utf-8")
    return d


def _bfcl_real_scores(monkeypatch, tmp_path, evaluate):
    """Harness stubs that keep the real _parse_scores / _clear_stale_scores."""
    monkeypatch.setenv("BFCL_PROJECT_ROOT", str(tmp_path))
    monkeypatch.setattr(BFCL, "check_bfcl_prerequisites",
                        lambda url: (True, "", list(BFCL.AGENTIC_CATEGORIES)))
    monkeypatch.setattr(BFCL, "search_backend", lambda: (BFCL.SEARCH_BACKEND, ""))
    monkeypatch.setattr(BFCL, "_register_model", lambda m: "org/gbench-model")

    def cli(argv, env, timeout, model=None, backend=None):
        return (0, "gen ok") if argv[0] == "generate" else evaluate(argv)

    monkeypatch.setattr(BFCL, "_run_cli", cli)


def test_bfcl_failed_evaluate_never_reports_an_earlier_runs_scores(monkeypatch, tmp_path):
    """The G3 review reproduction: stale 90% score files plus an evaluate that exits 1."""
    _write_scores(tmp_path, "org/gbench-model", BFCL.AGENTIC_CATEGORIES)
    _bfcl_real_scores(monkeypatch, tmp_path, lambda argv: (1, "Traceback: evaluate crashed"))
    # the harness registers the served model under its own id
    with pytest.raises(RuntimeError, match=r"bfcl evaluate` failed \(rc=1\)"):
        BFCL.run_bfcl_v4_agentic("org/gbench-model", "http://x", concurrency=1)
    left = list((tmp_path / "score").rglob("*_score.json"))
    assert left == [], "the stale files were removed before evaluate ran"


def test_bfcl_evaluate_that_writes_nothing_is_a_hard_error(monkeypatch, tmp_path):
    _write_scores(tmp_path, "org/gbench-model", BFCL.AGENTIC_CATEGORIES)
    _bfcl_real_scores(monkeypatch, tmp_path, lambda argv: (0, "done"))
    with pytest.raises(RuntimeError, match="wrote no score file"):
        BFCL.run_bfcl_v4_agentic("org/gbench-model", "http://x", concurrency=1)


def test_bfcl_reads_only_this_runs_selected_categories(monkeypatch, tmp_path):
    mem = list(BFCL.MEMORY_CATEGORIES)
    # leftovers: a web_search score of this model and another model's memory scores
    _write_scores(tmp_path, "org/gbench-model", ["web_search_base"], acc=1.0)
    _write_scores(tmp_path, "org/gbench-model-old", mem, acc=1.0)

    def evaluate(argv):
        _write_scores(tmp_path, "org/gbench-model", mem, acc=0.5)
        return 0, "done"

    _bfcl_real_scores(monkeypatch, tmp_path, evaluate)
    res = BFCL.run_bfcl_v4_agentic("org/gbench-model", "http://x", concurrency=1,
                                   eval_categories=",".join(mem))
    assert sorted(res["category_accuracy"]) == sorted(mem)
    assert res["accuracy"] == res["accuracy_pooled"] == 50.0
    assert res["summary_web_search"] is None


def test_bfcl_agentic_score_averages_unrounded_accuracies():
    third = {"correct": 1, "total": 3, "accuracy": 33.33}
    per_cat = {c: dict(third) for c in BFCL.MEMORY_CATEGORIES}
    per_cat.update({"web_search_base": {"correct": 2, "total": 3, "accuracy": 66.67},
                    "web_search_no_snippet": {"correct": 2, "total": 3, "accuracy": 66.67}})
    s = BFCL.agentic_scores(per_cat)
    assert s["summary_memory"] == 33.33 and s["summary_web_search"] == 66.67
    assert s["overall_accuracy_agentic"] == 50.0      # (33.33 + 66.67) / 2 from the exact values


# --------------------------------------------------------------------------- #
# gaia audio attachments are sized as audio, not as base64 text
# --------------------------------------------------------------------------- #
def _gaia_audio_messages(n_bytes=280_868):
    audio = GAIA._file_parts("1f975693.mp3", bytes(range(256)) * (n_bytes // 256))
    return [{"role": "system", "content": GAIA._GAIA_SYSTEM},
            {"role": "user", "content": [{"type": "text", "text": "What page numbers?"}] + audio}]


def test_gaia_audio_part_is_counted_as_soft_tokens_not_text():
    msgs = _gaia_audio_messages()
    b64 = msgs[1]["content"][2]["input_audio"]["data"]
    assert len(b64) > 370_000                          # the size of the largest GAIA mp3
    text, images, _n_messages, audios = base._prompt_text_and_images(msgs)
    assert audios == 1 and images == 0 and b64[:1000] not in text
    clean, _, n_audio = base._text_only_messages(msgs)
    assert n_audio == 1 and b64[:1000] not in json.dumps(clean)


def test_gaia_audio_task_is_neither_refused_nor_clamped_by_the_context_check():
    msgs = _gaia_audio_messages()
    budget, refused = base.clamp_to_context(msgs, 16384, max_model_len=131072)
    assert refused is None and budget == 16384


def test_gaia_audio_base64_is_truncated_in_traces():
    msgs = _gaia_audio_messages()
    part = base._sanitize_for_trace(msgs)[1]["content"][2]
    assert part["type"] == "input_audio" and part["input_audio"]["format"] == "mp3"
    assert len(part["input_audio"]["data"]) < 64


def test_docs_state_the_protocol_the_code_runs():
    gaia = " ".join(_doc("gaia").split())
    assert "GBENCH_AUDIO_SOFT_TOKENS" in gaia and "35.3 s" in gaia
    assert "20 tool-enabled model turns" in gaia
    dsqa = " ".join(_doc("deepsearch_qa").split())
    assert "Cost warning" in dsqa and "max_model_len" in dsqa
    bfcl = " ".join(_doc("bfcl_v4_agentic").split())
    assert "score files for the selected categories are deleted" in bfcl
    assert "--max-output-tokens" in bfcl and "needs the CLI to accept" not in bfcl
